"""Tests for export_mastr.py: select statements, naming and file conversion.

No database is involved: the select statements are built directly and the
CSV-to-Parquet/Excel conversion runs on a tiny generated CSV.
"""

import csv

import polars as pl
import pytest
from openpyxl import load_workbook

from db.entities import Einheiten, Marktakteur, Type
from db.polars_schema import build_polars_schema
from export_mastr import (
    ConvertExportJob,
    DataFormats,
    MastrExporter,
    MastrType,
    write_excel_parallel,
    write_parquet_parallel,
)

latinify = MastrExporter.latinify


def build_exporter() -> MastrExporter:
    """Build a MastrExporter without touching the database."""
    exporter = object.__new__(MastrExporter)
    exporter.states = {}
    exporter.csv_files_states = {k: [] for k in MastrType}
    exporter.csv_files_germany = {k: None for k in MastrType}
    exporter.select_stmts = {k: "" for k in MastrType}
    exporter._MastrExporter__create_select_stmts()
    return exporter


def output_names(statement):
    """Output column names: explicit alias if present, else the column name."""
    names = []
    for line in statement.splitlines():
        if not line.startswith("\t"):
            continue
        line = line.rstrip(",")
        if line.endswith(")") and " " in line:
            names.append(line.rsplit(" ", 1)[1])  # udf(...) Alias
        elif " " in line:
            names.append(line.rsplit(" ", 1)[1])  # table.column Alias
        else:
            names.append(line.split(".")[-1])  # bare table.column
    return names


class TestSelectStatements:
    @pytest.mark.parametrize("mastr_type", list(MastrType))
    def test_output_names_are_unique(self, mastr_type):
        names = output_names(build_exporter().select_stmts[mastr_type])
        assert len(names) == len(set(names)), "duplicate export column names"

    @pytest.mark.parametrize("mastr_type", list(MastrType))
    def test_output_names_match_polars_schema(self, mastr_type):
        """CSV/Excel headers and Parquet column names must be identical."""
        entity = mastr_type.value.entity
        names = output_names(build_exporter().select_stmts[mastr_type])
        assert names == list(build_polars_schema(entity).keys())

    def test_operator_columns_are_prefixed(self):
        statement = build_exporter().select_stmts[MastrType.WIND]
        for column in Marktakteur:
            assert f"Marktakteur_{column.name}" in statement
        assert "\tMarktakteure.Land Land,\n" not in statement
        assert "\tMarktakteure.Land,\n" not in statement

    def test_udfs_are_applied(self):
        statement = build_exporter().select_stmts[MastrType.WIND]
        assert "katalog(EinheitenWind.Land) Land," in statement
        assert "boolean_hr(EinheitenWind.FernsteuerbarkeitNb) FernsteuerbarkeitNb," in statement
        assert "katalog_list(EinheitenSolar.ZusaetzlicheMerkmaleDerFlaecheUndDerAktuellenFlaechennutzung)" in (
            build_exporter().select_stmts[MastrType.SOLAR]
        )

    def test_join_condition(self):
        for mastr_type in MastrType:
            statement = build_exporter().select_stmts[mastr_type]
            table = mastr_type.value.entity.name
            assert f"ON {table}.AnlagenbetreiberMastrNummer = Marktakteure.MastrNummer" in statement

    def test_exported_boolean_columns_are_udf_wrapped(self):
        """Regression: COPY renders raw booleans as t/f, which polars cannot
        parse as Boolean - every boolean column of an exported entity must
        carry a UDF (boolean_hr -> ja/nein)."""
        exported = [mastr_type.value.entity.value for mastr_type in MastrType] + [Marktakteur]
        for entity_class in exported:
            for column in entity_class:
                if column.value.type is Type.BOOLEAN:
                    assert column.value.udf is not None, (
                        f"{entity_class.__name__}.{column.name} is exported as raw boolean"
                    )


class TestLatinify:
    @pytest.mark.parametrize(
        "value,expected",
        [
            ("Baden-Württemberg", "Baden-Wuerttemberg"),
            ("Thüringen", "Thueringen"),
            ("Hessen", "Hessen"),
            ("Äpfel", "Aepfel"),
        ],
    )
    def test_umlauts(self, value, expected):
        assert latinify(value) == expected


def write_csv_with_schema(path, schema, rows):
    """Write a CSV matching the given polars schema (header + data rows)."""
    sample_values = {
        pl.String: "x",
        pl.Int64: "1",
        pl.Float64: "1.5",
        pl.Boolean: "true",
        pl.Datetime: "2024-01-01T00:00:00",
        pl.Date: "2024-01-01",
    }
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(schema.keys())
        for _ in range(rows):
            writer.writerow([sample_values[dtype] for dtype in schema.values()])


class TestFileConversion:
    def test_parquet_columns_match_schema(self, tmp_path):
        schema = build_polars_schema(Einheiten.EinheitenWind)
        csv_path = tmp_path / "wind.csv"
        parquet_path = tmp_path / "wind.parq"
        write_csv_with_schema(csv_path, schema, rows=2)

        job = ConvertExportJob(
            name="wind", file_output_path=parquet_path, csv_source_file=csv_path, schema=schema
        )
        write_parquet_parallel(job)

        frame = pl.read_parquet(parquet_path)
        assert frame.columns == list(schema.keys())
        assert frame.height == 2

    def test_excel_sheet_titles_fit_the_excel_limit(self):
        """Excel (all versions, including 365) rejects worksheet titles longer
        than 31 characters - openpyxl only warns, but Excel shows a repair
        dialog when opening such files. The file name carries the full
        identity, the sheet title stays short on purpose."""
        for mastr_type in MastrType:
            title = mastr_type.value.sheet_title
            assert len(title) <= 31, f"{mastr_type.name}: sheet title {title!r} exceeds Excel's 31-character limit"

    def test_large_datasets_have_no_excel_format(self):
        """Solar and Stromspeicher are too large for usable Excel files
        (openpyxl writes ~240k cells/s on a single core; Stromspeicher has
        ~1.5M units) - they are exported as CSV and Parquet only."""
        for mastr_type in (MastrType.SOLAR, MastrType.STROMSPEICHER):
            assert DataFormats.EXCEL not in mastr_type.value.data_formats, mastr_type.name

    def test_excel_contains_all_rows(self, tmp_path):
        schema = build_polars_schema(Einheiten.EinheitenSolar)
        csv_path = tmp_path / "solar.csv"
        excel_path = tmp_path / "solar.xlsx"
        write_csv_with_schema(csv_path, schema, rows=3)

        job = ConvertExportJob(name="solar", file_output_path=excel_path, csv_source_file=csv_path)
        write_excel_parallel(job)

        workbook = load_workbook(excel_path)
        sheet = workbook.active
        assert sheet.max_row == 4  # header + 3 data rows

    def test_excel_uses_the_sheet_title_not_the_job_name(self, tmp_path):
        """Regression: the sheet title used to be the job name (the file
        descriptor), which exceeds Excel's 31-character limit for many
        exports (e.g. geothermie_grubengas_druckentspannung_mecklenburg-vorpommern)."""
        import warnings

        schema = build_polars_schema(Einheiten.EinheitenWind)
        csv_path = tmp_path / "geothermie_grubengas_druckentspannung_mecklenburg-vorpommern.csv"
        excel_path = tmp_path / "geothermie_grubengas_druckentspannung_mecklenburg-vorpommern.xlsx"
        write_csv_with_schema(csv_path, schema, rows=2)

        job = ConvertExportJob(
            name="geothermie_grubengas_druckentspannung_mecklenburg-vorpommern",
            file_output_path=excel_path,
            csv_source_file=csv_path,
            sheet_title="Geothermie",
        )
        with warnings.catch_warnings():
            warnings.simplefilter("error")  # fail on the openpyxl UserWarning
            write_excel_parallel(job)

        workbook = load_workbook(excel_path)
        assert workbook.active.title == "Geothermie"
